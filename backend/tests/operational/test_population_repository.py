import shutil
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from app.agent.models import AgentExecutionContext
from app.agent.tools.operational_tools import (
    DeterministicKPITool,
    ReportLookupRequest,
    StoredValidationTool,
)
from app.models.operational_report import OperationalReportRecord
from app.operational_data.models import PopulationSourceRole
from app.operational_data.pipeline import (
    OperationalPipelineService,
    PopulationBundleProcessingError,
)
from app.operational_data.repository import SqlOperationalReportRepository
from app.reporting.repository import GeneratedReportRepository
from app.reporting.service import ReportGenerationService

DEMO = (
    Path(__file__).parents[3]
    / "ubnd_report_dataset"
    / "01_input_reports"
    / "population"
    / "normal"
    / "demo_case_1"
)
FILES = {
    PopulationSourceRole.OPENING_BALANCE: "01_so_du_dan_cu_dau_ky_2026-06.xlsx",
    PopulationSourceRole.CIVIL_STATUS: "02_bao_cao_ho_tich_thang_06_2026.xlsx",
    PopulationSourceRole.RESIDENCE_MOVEMENT: "03_bao_cao_bien_dong_cu_tru_06_2026.xlsx",
}


def source_mapping(overrides=None):
    overrides = overrides or {}
    return {
        role: (overrides.get(role, DEMO / filename), filename)
        for role, filename in FILES.items()
    }


def test_population_bundle_persistence_is_idempotent(db_session, tmp_path):
    pipeline = OperationalPipelineService(
        SqlOperationalReportRepository(db_session), tmp_path / "operational", 10
    )
    record, duplicate = pipeline.ingest_population_bundle(source_mapping())
    repeated, repeated_duplicate = pipeline.ingest_population_bundle(source_mapping())

    assert not duplicate
    assert repeated_duplicate
    assert repeated.id == record.id
    assert record.values["population_closing"] == 12458
    assert record.values["temporary_closing"] == 1174
    assert db_session.query(OperationalReportRecord).count() == 1
    assert len(list(Path(record.file_path).glob("*.xlsx"))) == 3

    request = ReportLookupRequest(report_id=record.id)
    context = AgentExecutionContext(request_id="population-canonical")
    kpis = DeterministicKPITool(db_session).execute(request, context)
    validation = StoredValidationTool(db_session).execute(request, context)
    assert kpis.data["deterministic_kpis"]["population_net_change"] == "28"
    assert kpis.data["deterministic_kpis"]["temporary_net_change"] == "54"
    assert validation.data["ruleset"] == "POPULATION_CANONICAL_DETERMINISTIC_V1"

    generated = ReportGenerationService(
        GeneratedReportRepository(db_session), tmp_path / "reports"
    ).generate(record, "synthetic_domain_draft")
    document = Document(generated.artifact_path)
    rendered_cells = [
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    ]
    assert "Tổng khai sinh đăng ký" in rendered_cells
    assert "Khai sinh thuộc dân cư thường trú của xã" in rendered_cells
    assert "Tổng khai tử đăng ký" in rendered_cells
    assert "Khai tử thuộc dân cư thường trú của xã" in rendered_cells


def test_different_bundle_for_same_business_key_requires_confirmation(
    db_session, tmp_path
):
    pipeline = OperationalPipelineService(
        SqlOperationalReportRepository(db_session), tmp_path / "operational", 10
    )
    pipeline.ingest_population_bundle(source_mapping())
    changed = tmp_path / FILES[PopulationSourceRole.OPENING_BALANCE]
    shutil.copy2(DEMO / FILES[PopulationSourceRole.OPENING_BALANCE], changed)
    workbook = load_workbook(changed)
    workbook["Dinh_nghia_chi_tieu"]["F7"] = "Synthetic checksum variation"
    workbook.save(changed)
    workbook.close()

    with pytest.raises(PopulationBundleProcessingError) as captured:
        pipeline.ingest_population_bundle(
            source_mapping({PopulationSourceRole.OPENING_BALANCE: changed})
        )
    assert captured.value.issues[0].code == "POP_BUSINESS_KEY_CONFLICT"
    assert db_session.query(OperationalReportRecord).count() == 1
